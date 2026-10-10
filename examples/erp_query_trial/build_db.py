"""Import the five ERP exports into an isolated, traceable SQLite database."""

import argparse
from collections import defaultdict
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
import json
from pathlib import Path
import sqlite3
from typing import Any

from openpyxl import load_workbook


FILES = {
    "production_receipts": "2025年生产入库单列表.xlsx",
    "customers": "客户_2026092215331494_100137.xlsx",
    "materials": "物料_2026092215325455_100137.xlsx",
    "sales_outbound": "销售出库单_2026092215442520_100137.xlsx",
    "sales_orders": "销售订单_2026092215154329_100137(1).xlsx",
}


def text(value: Any) -> str | None:
    if value is None or value == "":
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def number(value: Any) -> float | None:
    value = text(value)
    if value is None:
        return None
    parsed = Decimal(value)
    if not parsed.is_finite():
        raise ValueError("Non-finite number")
    return float(parsed)


def iso_date(value: Any, include_time: bool = False) -> str | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        parsed = value
    elif isinstance(value, date):
        parsed = datetime.combine(value, datetime.min.time())
    else:
        normalized = str(value).replace("/", "-")
        try:
            parsed = datetime.fromisoformat(normalized)
        except ValueError:
            parsed = datetime.strptime(normalized, "%Y-%m-%d")
    return parsed.isoformat(sep=" ") if include_time else parsed.date().isoformat()


def insert(conn: sqlite3.Connection, table: str, values: dict) -> None:
    # Table/column names come exclusively from this module, not input workbooks.
    columns = ",".join(values)
    placeholders = ",".join("?" for _ in values)
    conn.execute(
        f"INSERT INTO {table} ({columns}) VALUES ({placeholders})",
        tuple(values.values()),
    )


def unpack_rows(keys: list, rows: list, first_row: int) -> list[tuple[int, dict]]:
    if len(keys) != len(set(keys)):
        raise ValueError("Duplicate source column keys")
    return [
        (index, {**dict.fromkeys(keys), **dict(zip(keys, row))})
        for index, row in enumerate(rows, first_row)
    ]


def bill_entries(
    keys: list, rows: list, entry_key: str, first_row: int = 3
) -> list[tuple[int, dict, dict]]:
    """Carry only header/finance fields inside an explicitly started document."""
    head_key = keys[0]
    header_end = keys.index("*Split*2")
    current: dict = {}
    result = []
    for row_number, row in unpack_rows(keys, rows, first_row):
        if text(row.get(head_key)) is not None:
            current = {k: row.get(k) for k in keys[:header_end]}
        elif not current:
            raise ValueError(f"Row {row_number}: continuation without document")
        if text(row.get(entry_key)) is not None:
            result.append((row_number, current.copy(), row))
    return result


def add_issue(conn: sqlite3.Connection, kind: str, table: str, row: dict,
              detail: str) -> None:
    insert(conn, "data_issues", {
        "issue_type": kind, "business_table": table,
        "sheet_id": row.get("sheet_id"), "excel_row": row.get("excel_row"),
        "detail": detail,
    })


def production_entries(keys: list, rows: list) -> list[tuple[int, dict]]:
    current = {}
    result = []
    for row_number, row in unpack_rows(keys, rows, 2):
        if text(row.get("单据编号")) is not None:
            current = {key: row[key] for key in keys[:5]}
        elif not current:
            raise ValueError(f"Row {row_number}: production continuation without bill")
        combined = row.copy()
        for key in keys[:5]:
            if combined[key] is None:
                combined[key] = current[key]
        result.append((row_number, combined))
    return result


def import_sheet(conn: sqlite3.Connection, path: Path, ws: Any, cached_ws: Any,
                 header_count: int) -> tuple[int, list, list]:
    # Some ERP exports incorrectly declare A1:A1 as their used range.
    ws.reset_dimensions()
    cached_ws.reset_dimensions()
    cells = list(ws.iter_rows())
    values = list(cached_ws.iter_rows(values_only=True))
    if len(cells) != len(values):
        raise ValueError("Formula and cached worksheet row counts differ")
    formulas = []
    for row_index, row in enumerate(cells):
        formula_row = {}
        for column_index, cell in enumerate(row):
            if cell.data_type == "f":
                if values[row_index][column_index] is None:
                    raise ValueError(f"{ws.title}/{cell.coordinate}: missing formula cache")
                formula_row[str(column_index + 1)] = cell.value
        formulas.append(formula_row)
    keys = values[0]
    labels = values[1] if header_count == 2 else keys
    if any(k is None for k in keys) or len(keys) != len(set(keys)):
        raise ValueError(f"{path.name}/{ws.title}: missing/duplicate column keys")
    cursor = conn.execute(
        "INSERT INTO source_sheets(workbook,sheet,header_rows) VALUES (?,?,?)",
        (path.name, ws.title, header_count),
    )
    sheet_id = cursor.lastrowid
    for i, (key, label) in enumerate(zip(keys, labels), 1):
        conn.execute("INSERT INTO source_columns VALUES (?,?,?,?)",
                     (sheet_id, i, text(key), text(label)))
    rows = values[header_count:]
    for row_number, row in enumerate(rows, header_count + 1):
        conn.execute("INSERT INTO raw_rows VALUES (?,?,?,?)", (
            sheet_id, row_number,
            json.dumps(row, ensure_ascii=False, default=lambda v: v.isoformat()),
            json.dumps(formulas[row_number - 1], ensure_ascii=False),
        ))
    return sheet_id, keys, rows


def import_master(conn: sqlite3.Connection, table: str, sheet_id: int,
                  keys: list, rows: list) -> None:
    for row_number, row in unpack_rows(keys, rows, 2):
        common = {"sheet_id": sheet_id, "excel_row": row_number}
        if table == "customers":
            common.update(
                customer_code=text(row["客户编码"]),
                customer_name=text(row["客户名称"]),
                short_name=text(row["简称"]),
                organization=text(row["使用组织"]),
                customer_group=text(row["客户分组"]),
                document_status=text(row["单据状态"]), disabled=text(row["禁用状态"]),
            )
        else:
            common.update(
                material_code=text(row["编码"]), material_name=text(row["名称"]),
                specification=text(row["规格型号"]),
                organization=text(row["使用组织"]), base_unit=text(row["基本单位"]),
                material_group=text(row["物料分组"]),
                inventory_category=text(row["存货类别"]),
                document_status=text(row["数据状态"]), disabled=text(row["禁用状态"]),
            )
        insert(conn, table, common)


def import_sales(conn: sqlite3.Connection, table: str, sheet_id: int,
                 keys: list, rows: list) -> None:
    order = table == "sales_orders"
    entry_key = "FSaleOrderEntry" if order else "FEntity"
    customer_key = "FCustId" if order else "FCustomerID"
    material_key = "FMaterialId" if order else "FMaterialID"
    currency_key = "FSettleCurrId" if order else "FSettleCurrID"
    for row_number, head, row in bill_entries(keys, rows, entry_key):
        record = {
            "entry_id": text(row[entry_key]), "document_id": text(head[keys[0]]),
            "bill_no": text(head["FBillNo"]), "bill_date": iso_date(head["FDate"]),
            "customer_code": text(head[customer_key]),
            "customer_name": text(head[customer_key + "#Name"]),
            "sales_org_code": text(head["FSaleOrgId"]),
            "currency_code": text(head[currency_key]),
            "material_code": text(row[material_key]),
            "material_name": text(row[material_key + "#Name"]),
            "customer_material_code": text(row["FMapId" if order else "FCustMatID"]),
            "length_value": text(row["FAuxPropId#FF100001"]),
            "brand_code": text(row["FAuxPropId#FF100002"]),
            "brand_name": text(row["FAuxPropId#FF100002#Name"]),
            "unit": text(row["FUnitID"]),
            "quantity": number(row["FQty" if order else "FRealQty"]),
            "unit_price": number(row["FPrice"]),
            "tax_unit_price": number(row["FTaxPrice"]),
            "current_inventory": number(row["FInventoryQty"]),
            "document_status": None,
            "sheet_id": sheet_id, "excel_row": row_number,
        }
        if order:
            record.update(
                requested_delivery_at=iso_date(row["FDeliveryDate"], True),
                available_inventory=number(row["FCurrentInventory"]),
                return_type=text(row["FReturnType"]),
            )
        else:
            record.update(
                warehouse_code=text(row["FStockID"]),
                warehouse_name=text(row["FStockID#Name"]),
                source_order_entry_id=text(row["FSOEntryId"]),
            )
        insert(conn, table, record)


def import_production(conn: sqlite3.Connection, table: str, sheet_id: int,
                      keys: list, rows: list) -> None:
    names = defaultdict(list)
    for code, name in conn.execute("SELECT customer_code,customer_name FROM customers"):
        names[name].append(code)
    records = (production_entries(keys, rows) if table == "production_receipts"
               else unpack_rows(keys, rows, 2))
    for row_number, row in records:
        common = {"sheet_id": sheet_id, "excel_row": row_number,
                  "material_code": text(row["物料编码"]),
                  "material_name": text(row["物料名称"])}
        if table == "production_standards":
            if common["material_code"] is None:
                add_issue(conn, "missing_material_code", table, common,
                          "原始工时行没有物料编码，仅保留原始行，未生成工时记录")
                continue
            numbers = {}
            for field, key in (("assembly_value", "组装"), ("coating_value", "镀膜"),
                               ("total_labor_cost", "总工费"),
                               ("total_labor_minutes", "总工时-分钟")):
                try:
                    numbers[field] = number(row[key])
                except (InvalidOperation, ValueError):
                    numbers[field] = None
                    add_issue(conn, "invalid_numeric", table, common,
                              f"{key}: {text(row[key])}")
            common.update(
                **numbers,
            )
        else:
            customer_name = text(row["客户"])
            matches = names[customer_name]
            common.update(
                bill_no=text(row["单据编号"]), bill_date=iso_date(row["日期"]),
                customer_name=customer_name,
                customer_code=matches[0] if len(matches) == 1 else None,
                terminal_customer_name=text(row["终端客户"]),
                product_type=text(row["产品类型"]), unit=text(row["单位"]),
                quantity=number(row["实收数量"]), warehouse_name=text(row["仓库"]),
                length_value=text(row["辅助属性.长度"]),
                brand_code=text(row["辅助属性.品牌.编码"]),
                brand_name=text(row["辅助属性.品牌.名称"]),
                document_status=text(row["单据状态"]),
            )
            if customer_name and len(matches) != 1:
                add_issue(conn, "customer_name_unresolved", table, common,
                          "客户名称无法唯一匹配客户主档；未进行模糊匹配")
        insert(conn, table, common)


def check_links(conn: sqlite3.Connection) -> None:
    for table in ("sales_orders", "sales_outbound", "production_receipts",
                  "production_standards"):
        for key, master in (("material_code", "materials"),
                            ("customer_code", "customers")):
            if table == "production_standards" and key == "customer_code":
                continue
            unmatched = conn.execute(
                f"SELECT t.sheet_id,t.excel_row,t.{key} FROM {table} t "
                f"LEFT JOIN {master} m ON t.{key}=m.{key} "
                f"WHERE t.{key} IS NOT NULL AND m.{key} IS NULL"
            ).fetchall()
            for sheet_id, excel_row, code in unmatched:
                add_issue(conn, "master_missing_" + key, table,
                          {"sheet_id": sheet_id, "excel_row": excel_row}, code)
    for row in conn.execute(
        "SELECT s.sheet_id,s.excel_row,s.source_order_entry_id "
        "FROM sales_outbound s LEFT JOIN sales_orders o "
        "ON o.entry_id=s.source_order_entry_id WHERE o.entry_id IS NULL"
    ).fetchall():
        add_issue(conn, "order_link_not_in_export", "sales_outbound",
                  {"sheet_id": row[0], "excel_row": row[1]},
                  row[2] or "源订单明细ID缺失")
    for row in conn.execute(
        "SELECT s.sheet_id,s.excel_row FROM sales_outbound s "
        "JOIN sales_orders o ON o.entry_id=s.source_order_entry_id "
        "WHERE s.material_code<>o.material_code OR s.unit<>o.unit "
        "OR s.customer_code<>o.customer_code "
        "OR s.length_value IS NOT o.length_value "
        "OR s.brand_code IS NOT o.brand_code"
    ).fetchall():
        add_issue(conn, "order_link_attribute_mismatch", "sales_outbound",
                  {"sheet_id": row[0], "excel_row": row[1]},
                  "源订单关联存在，但客户、物料、单位或辅助属性不一致")


def build(source_dir: Path, output: Path) -> dict:
    if output.exists():
        raise FileExistsError(f"Refusing to overwrite {output}; choose a new output")
    for name in FILES.values():
        if not (source_dir / name).is_file():
            raise FileNotFoundError(name)
    output.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(output)
    try:
        conn.executescript(Path(__file__).with_name("schema.sql").read_text("utf-8"))
        with conn:
            for table in ("customers", "materials", "sales_orders",
                          "sales_outbound", "production_receipts"):
                path = source_dir / FILES[table]
                wb = load_workbook(path, read_only=True, data_only=False)
                cached_wb = load_workbook(path, read_only=True, data_only=True)
                try:
                    if len(wb.worksheets) != (2 if table == "production_receipts" else 1):
                        raise ValueError(f"Unexpected worksheet count: {path.name}")
                    for i, ws in enumerate(wb.worksheets):
                        target = "production_standards" if i == 1 else table
                        sheet_id, keys, rows = import_sheet(
                            conn, path, ws, cached_wb[ws.title],
                            2 if table.startswith("sales_") else 1
                        )
                        if table in ("customers", "materials"):
                            import_master(conn, target, sheet_id, keys, rows)
                        elif table.startswith("sales_"):
                            import_sales(conn, target, sheet_id, keys, rows)
                        else:
                            import_production(conn, target, sheet_id, keys, rows)
                finally:
                    wb.close()
                    cached_wb.close()
            check_links(conn)
            if conn.execute("PRAGMA foreign_key_check").fetchall():
                raise ValueError("Broken source provenance")
        report = {
            "created_at_utc": datetime.now(timezone.utc).isoformat(),
            "database": output.name,
            "counts": {table: conn.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0]
                       for table in (*FILES, "production_standards", "raw_rows")},
            "date_ranges": {table: conn.execute(
                f"SELECT MIN(bill_date),MAX(bill_date) FROM {table}"
            ).fetchone() for table in
                ("sales_orders", "sales_outbound", "production_receipts")},
            "issues": [dict(zip(("type", "table", "count"), row)) for row in
                       conn.execute("SELECT * FROM v_data_issue_summary")],
            "limitations": [
                "静态导出样本，不是实时ERP；未连接正式数据库",
                "销售订单/出库导出未提供审核与作废状态，不能判定有效单据",
                "库存字段为空，没有退货表，不能计算当前库存/净销售/真实欠交",
                "订单、出库与生产入库时间范围不同，关联缺失不一定是错误",
                "金额未推算；数量按单位分别统计；长度、品牌保留为文本属性",
                "收款、交货计划等次级实体保留在原始行，未混入产品明细",
            ],
        }
        return report
    except BaseException:
        conn.close()
        output.unlink(missing_ok=True)
        raise
    finally:
        conn.close()


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source-dir", type=Path, required=True)
    parser.add_argument("--output", type=Path,
                        default=Path(__file__).parent / "data" / "erp_trial.sqlite")
    args = parser.parse_args()
    report = build(args.source_dir, args.output)
    report_path = args.output.with_suffix(".import.json")
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), "utf-8")
    print(json.dumps(report, ensure_ascii=True, indent=2))


if __name__ == "__main__":
    main()
