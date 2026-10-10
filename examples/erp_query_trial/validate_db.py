"""Reconcile the SQLite database against the original Excel exports."""

import argparse
import json
import re
from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.cell import column_index_from_string
from query_db import connect_readonly, report


def identifier(value):
    return (
        str(int(value))
        if isinstance(value, float) and value.is_integer()
        else str(value)
    )


def normalized_date(value):
    if isinstance(value, (date, datetime)):
        return value.strftime("%Y-%m-%d")
    return datetime.strptime(
        str(value).split(" ")[0].replace("/", "-"), "%Y-%m-%d"
    ).strftime("%Y-%m-%d")


def check(conn, source_dir: Path) -> dict:
    checks = []
    formula_count = 0
    raw_count = 0
    expected_totals = defaultdict(lambda: Decimal(0))
    normalized_rows = 0
    independent_ranking = defaultdict(lambda: Decimal(0))
    for sheet_id, workbook, sheet_name, headers in conn.execute(
        "SELECT id,workbook,sheet,header_rows FROM source_sheets"
    ):
        source = source_dir / workbook
        wb = load_workbook(source, read_only=True, data_only=True)
        formulas_wb = load_workbook(source, read_only=True, data_only=False)
        try:
            ws = wb[sheet_name]
            fws = formulas_wb[sheet_name]
            ws.reset_dimensions()
            fws.reset_dimensions()
            rows = list(ws.iter_rows(values_only=True))
            formula_rows = list(fws.iter_rows())
            stored = conn.execute(
                "SELECT excel_row,values_json,formulas_json FROM raw_rows "
                "WHERE sheet_id=? ORDER BY excel_row",
                (sheet_id,),
            ).fetchall()
            assert len(stored) == len(rows) - headers, (workbook, "raw count")
            for row_number, values_json, formulas_json in stored:
                raw_count += 1
                source_values = json.loads(
                    json.dumps(
                        rows[row_number - 1],
                        ensure_ascii=False,
                        default=lambda v: v.isoformat(),
                    )
                )
                assert json.loads(values_json) == source_values, (workbook, row_number)
                actual_formulas = {
                    str(i + 1): cell.value
                    for i, cell in enumerate(formula_rows[row_number - 1])
                    if cell.data_type == "f"
                }
                assert json.loads(formulas_json) == actual_formulas
                for column, formula in actual_formulas.items():
                    formula_count += 1
                    direct = re.fullmatch(r"=([A-Z]+)(\d+)", formula)
                    labor = re.fullmatch(r"=E(\d+)\*0\.6/24\*60", formula)
                    cached = source_values[int(column) - 1]
                    if direct:
                        expected = rows[int(direct[2]) - 1][
                            column_index_from_string(direct[1]) - 1
                        ]
                        assert cached == expected, (workbook, row_number, column)
                    elif labor:
                        cost = rows[int(labor[1]) - 1][4]
                        if isinstance(cost, str) and cost.startswith("#"):
                            assert cached == cost
                        else:
                            computed = Decimal(str(cost)) * Decimal("0.6") / 24 * 60
                            assert abs(Decimal(str(cached)) - computed) < Decimal(
                                "0.000001"
                            )
                    else:
                        raise AssertionError(f"Unverified formula: {formula}")
            if workbook.startswith("客户_"):
                table, code_index = "customers", 0
            elif workbook.startswith("物料_"):
                table, code_index = "materials", 1
            elif workbook.startswith("销售订单_"):
                table, code_index = "sales_orders", 101
            elif workbook.startswith("销售出库单_"):
                table, code_index = "sales_outbound", 112
            else:
                table = (
                    "production_receipts"
                    if sheet_name == "Sheet1"
                    else "production_standards"
                )
                code_index = 5 if table == "production_receipts" else 0
            business = {
                r["excel_row"]: dict(r)
                for r in conn.execute(
                    f"SELECT * FROM {table} WHERE sheet_id=?", (sheet_id,)
                )
            }
            current = None
            expected_rows = set()
            for row_number, raw in enumerate(rows[headers:], headers + 1):
                values = list(raw) + [None] * (len(rows[0]) - len(raw))
                if table in ("sales_orders", "sales_outbound"):
                    if values[0] is not None:
                        current = values
                    entry_index = 92 if table == "sales_orders" else 108
                    if values[entry_index] is None:
                        continue
                elif table == "production_receipts":
                    if values[0] is not None:
                        current = values
                elif table == "production_standards" and values[0] is None:
                    continue
                expected_rows.add(row_number)
                record = business[row_number]
                key = "customer_code" if table == "customers" else "material_code"
                assert record[key] == identifier(values[code_index])
                normalized_rows += 1
                if table in ("sales_orders", "sales_outbound", "production_receipts"):
                    assert current is not None
                    date_index = 1 if table == "production_receipts" else 4
                    assert record["bill_date"] == normalized_date(current[date_index])
                    assert record["bill_no"] == identifier(
                        current[0 if table == "production_receipts" else 3]
                    )
                    if table != "production_receipts":
                        assert record["entry_id"] == identifier(values[entry_index])
                        assert record["document_id"] == identifier(current[0])
                        assert record["customer_code"] == identifier(current[7])
                    unit_index, qty_index = {
                        "sales_orders": (110, 116),
                        "sales_outbound": (117, 122),
                        "production_receipts": (8, 9),
                    }[table]
                    assert record["unit"] == values[unit_index]
                    quantity = Decimal(str(values[qty_index]))
                    assert Decimal(str(record["quantity"])) == quantity
                    expected_totals[(table, record["unit"])] += quantity
                    if (
                        table == "sales_outbound"
                        and record["unit"] == "Pcs"
                        and ("2025-08-18" <= record["bill_date"] < "2025-08-25")
                    ):
                        rank_key = (
                            identifier(values[112]),
                            values[113],
                            identifier(values[114])
                            if values[114] is not None
                            else None,
                            identifier(values[115])
                            if values[115] is not None
                            else None,
                            values[116],
                            values[117],
                        )
                        independent_ranking[rank_key] += quantity
            assert set(business) == expected_rows, (table, "row coverage")
            checks.append(
                {
                    "workbook": workbook,
                    "sheet": sheet_name,
                    "raw_rows": len(stored),
                    "business_rows": len(business),
                }
            )
        finally:
            wb.close()
            formulas_wb.close()
    totals = []
    for (table, unit), expected in expected_totals.items():
        actual = conn.execute(
            f"SELECT SUM(quantity) FROM {table} WHERE unit=?", (unit,)
        ).fetchone()[0]
        assert abs(Decimal(str(actual)) - expected) < Decimal("0.000001")
        totals.append(
            {
                "table": table,
                "unit": unit,
                "source_total": str(expected),
                "database_total": actual,
            }
        )
    ranked = report(conn, "shipments", "2025-08-18", "2025-08-25", "Pcs")
    expected_top = sorted(
        independent_ranking.items(),
        key=lambda p: (-p[1], p[0][0], p[0][2] or "", p[0][3] or ""),
    )[:10]
    assert len(ranked) == len(expected_top)
    for actual, (key, quantity) in zip(ranked, expected_top):
        assert (
            tuple(
                actual[c]
                for c in [
                    "material_code",
                    "material_name",
                    "length_value",
                    "brand_code",
                    "brand_name",
                    "unit",
                ]
            )
            == key
        )
        assert Decimal(str(actual["shipped_quantity"])) == quantity
    assert conn.execute("PRAGMA integrity_check").fetchone()[0] == "ok"
    assert not conn.execute("PRAGMA foreign_key_check").fetchall()
    assert (
        conn.execute("SELECT COUNT(*) FROM v_sales_order_shipments").fetchone()[0]
        == (conn.execute("SELECT COUNT(*) FROM sales_orders").fetchone()[0])
    )
    return {
        "result": "PASS",
        "source_checks": checks,
        "raw_rows_verified": raw_count,
        "business_rows_verified": normalized_rows,
        "formula_caches_verified": formula_count,
        "quantity_totals": totals,
        "sample_ranking": "2025-08-18至2025-08-24，Pcs，前10名与源表独立计算一致",
        "erp_reconciliation": "尚未与ERP在线报表对账",
    }


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source-dir", type=Path, required=True)
    parser.add_argument(
        "--db", type=Path, default=Path(__file__).parent / "data" / "erp_trial.sqlite"
    )
    args = parser.parse_args()
    conn = connect_readonly(args.db)
    try:
        result = check(conn, args.source_dir)
    finally:
        conn.close()
    args.db.with_suffix(".validation.json").write_text(
        json.dumps(result, ensure_ascii=False, indent=2), "utf-8"
    )
    print(json.dumps(result, ensure_ascii=True, indent=2))


if __name__ == "__main__":
    main()
